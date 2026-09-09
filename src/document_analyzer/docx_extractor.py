import os
import sys

from docx import Document as DocxDocument
from langchain_core.documents import Document

from exception import DocumentPortalException
from logger import get_logger
from src.document_analyzer.image_extractor import caption_image

log = get_logger(__name__)


def extract_docx_tables(file_path: str):
    """
    Extract tables from a DOCX file -- same design as PDF table extraction:
    one row = one Document, first row treated as headers.
    """
    if not os.path.exists(file_path):
        log.error(f"File does not exist: {file_path}")
        raise DocumentPortalException(f"File does not exist: {file_path}", sys)

    try:
        log.info(f"Extracting tables from DOCX: {file_path}")
        docx_doc = DocxDocument(file_path)
        docs = []

        for table_idx, table in enumerate(docx_doc.tables):
            if len(table.rows) < 1:
                continue

            headers = [
                cell.text.strip() or f"col_{i}" for i, cell in enumerate(table.rows[0].cells)
            ]
            data_rows = table.rows[1:] if len(table.rows) > 1 else []

            for row_idx, row in enumerate(data_rows):
                row_text = "\n".join(
                    f"{headers[i]}: {cell.text.strip()}" for i, cell in enumerate(row.cells)
                )
                if not row_text.strip():
                    continue

                docs.append(
                    Document(
                        page_content=row_text,
                        metadata={
                            "source": file_path,
                            "type": "table_row",
                            "table_index": table_idx,
                            "row_index": row_idx,
                        },
                    )
                )

        log.info(f"Extracted {len(docs)} table rows from DOCX: {file_path}")
        return docs

    except DocumentPortalException:
        raise
    except Exception as e:
        raise DocumentPortalException(
            f"Failed to extract tables from DOCX: {file_path}", sys
        ) from e


def extract_docx_images(file_path: str, caption: bool = True):
    """
    Extract images from a DOCX file -- same design as PDF image extraction:
    one image = one Document, page_content is an LLM-generated caption.
    """
    if not os.path.exists(file_path):
        log.error(f"File does not exist: {file_path}")
        raise DocumentPortalException(f"File does not exist: {file_path}", sys)

    try:
        log.info(f"Extracting images from DOCX: {file_path}")
        docx_doc = DocxDocument(file_path)
        docs = []

        image_parts = [
            part for part in docx_doc.part.related_parts.values() if "image" in part.content_type
        ]

        for img_idx, part in enumerate(image_parts):
            image_bytes = part.blob
            image_ext = part.content_type.split("/")[-1]

            if caption:
                caption_text = caption_image(image_bytes, image_ext)
            else:
                caption_text = f"[Uncaptioned image {img_idx}]"

            docs.append(
                Document(
                    page_content=caption_text,
                    metadata={
                        "source": file_path,
                        "type": "image",
                        "image_index": img_idx,
                        "image_format": image_ext,
                    },
                )
            )

        log.info(f"Extracted {len(docs)} images from DOCX: {file_path}")
        return docs

    except DocumentPortalException:
        raise
    except Exception as e:
        raise DocumentPortalException(
            f"Failed to extract images from DOCX: {file_path}", sys
        ) from e
